import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]

COLOR_RED = RGBColor(198, 40, 40)       # #C62828
COLOR_ORANGE = RGBColor(230, 81, 0)     # #E65100
COLOR_BLACK = RGBColor(17, 24, 39)      # #111827
COLOR_GRAY = RGBColor(75, 85, 99)       # #4B5563
COLOR_GREEN = RGBColor(46, 125, 50)     # #2E7D32
FONT_FAMILY = "Century Gothic"

def add_header(slide, title, subtitle=None):
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_FAMILY
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    
    if subtitle:
        run = p.add_run()
        run.text = f"   —   {subtitle}"
        run.font.name = FONT_FAMILY
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_ORANGE

def add_footer(slide, current, total=28):
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.3))
    tf = txBox.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "KOINEXA  |  Internal DSA Architecture & System Engineering"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_GRAY
    
    run = p.add_run()
    run.text = f"                                                                                                                    {current} / {total}"
    run.font.name = FONT_FAMILY
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = COLOR_ORANGE

print("Generator skeleton ready")
