import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

COLOR_RED = RGBColor(198, 40, 40)       # #C62828
COLOR_ORANGE = RGBColor(230, 81, 0)     # #E65100
COLOR_BLACK = RGBColor(17, 24, 39)      # #111827
COLOR_GRAY = RGBColor(75, 85, 99)       # #4B5563
COLOR_GREEN = RGBColor(46, 125, 50)     # #2E7D32
FONT_FAMILY = "Century Gothic"

def add_header(slide, title, subtitle=None):
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.75))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_FAMILY
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    if subtitle:
        run = p.add_run()
        run.text = f"   —   {subtitle}"
        run.font.name = FONT_FAMILY
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_ORANGE

def add_footer(slide, current, total=28):
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.3))
    tf = tx.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "KOINEXA  |  Internal DSA Architecture & System Engineering"
    p.font.name = FONT_FAMILY
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_GRAY
    run = p.add_run()
    run.text = f"                                                                                                                    {current} / {total}"
    run.font.name = FONT_FAMILY
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = COLOR_ORANGE

def add_col_section(tf, heading, color, bullets, space_after=12):
    p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
    p.text = heading
    p.font.name = FONT_FAMILY
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = color
    p.space_after = Pt(4)
    
    for b in bullets:
        bp = tf.add_paragraph()
        bp.space_after = Pt(3)
        bp.level = 0
        if isinstance(b, tuple):
            prefix, text = b
            r1 = bp.add_run()
            r1.text = prefix + ": "
            r1.font.name = FONT_FAMILY
            r1.font.size = Pt(9.5)
            r1.font.bold = True
            r1.font.color.rgb = COLOR_BLACK
            
            r2 = bp.add_run()
            r2.text = text
            r2.font.name = FONT_FAMILY
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = COLOR_BLACK
        else:
            r = bp.add_run()
            r.text = b
            r.font.name = FONT_FAMILY
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_BLACK
    
    if space_after:
        sp = tf.add_paragraph()
        sp.space_after = Pt(space_after)

def add_code_block(tf, heading, color, code_lines):
    p = tf.add_paragraph()
    p.text = heading
    p.font.name = FONT_FAMILY
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = color
    p.space_after = Pt(3)
    
    cp = tf.add_paragraph()
    cp.space_after = Pt(8)
    for line in code_lines:
        r = cp.add_run()
        r.text = line + "\n"
        r.font.name = "Courier New"
        r.font.size = Pt(8.0)
        r.font.color.rgb = COLOR_BLACK

print("Helper functions ready")
